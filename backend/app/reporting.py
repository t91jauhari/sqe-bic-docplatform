"""Generate reports from frozen snapshots, never from a generative decision."""

import io
import json
from html import escape

from openpyxl import Workbook
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Image


def csv_safe(value):
    value = (
        ""
        if value is None
        else value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False)
    )
    # Neutralize spreadsheet formulas in untrusted OCR text.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def make_report(report, fmt):
    snapshot = report.snapshot
    if fmt == "json":
        return json.dumps(
            {
                "version": report.version,
                "state": report.state,
                "approved_by": report.approved_by,
                **snapshot,
            },
            ensure_ascii=False,
            indent=2,
        ).encode(), "application/json"
    if fmt == "xlsx" and "scores" in snapshot:
        from .kertas_kerja import workbook

        return workbook(
            [
                {
                    **snapshot,
                    "release_state": report.state,
                    "snapshot_hash": report.fingerprint,
                }
            ]
        ), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Kertas Kerja"
        for row in [
            ["Report", snapshot["report_title"]],
            ["Case", snapshot["reference"]],
            ["Customer", snapshot["customer"]],
            ["Organization", snapshot["organization"]],
            ["Version", report.version],
            ["Release state", report.state],
            ["Configuration version", snapshot["config_version"]],
            ["Provider", snapshot["provider"]],
            [],
        ]:
            sheet.append([csv_safe(x) if isinstance(x, str) else x for x in row])
        sheet.append(
            ["Rule", "Field", "Result", "Message", "Actual", "Expected", "Evidence"]
        )
        for f in snapshot["findings"]:
            sheet.append(
                [
                    csv_safe(f.get(x))
                    for x in [
                        "rule_id",
                        "field",
                        "result",
                        "message",
                        "actual",
                        "expected",
                        "evidence",
                    ]
                ]
            )
        sheet.freeze_panes = "A11"
        for col in "ABCDEFG":
            sheet.column_dimensions[col].width = 30
        docs = book.create_sheet("Documents")
        docs.append(["Filename", "Type", "Status", "SHA256"])
        for d in snapshot["documents"]:
            docs.append(
                [csv_safe(d[x]) for x in ["filename", "doc_type", "status", "sha256"]]
            )
        fields = book.create_sheet("Extracted fields")
        fields.append(
            [
                "Document",
                "Type",
                "Field",
                "Raw",
                "Effective normalized",
                "Review override",
            ]
        )
        for d in snapshot["documents"]:
            for key, value in d["normalized"].items():
                fields.append(
                    [
                        csv_safe(d["filename"]),
                        d["doc_type"],
                        key,
                        csv_safe(d["raw"].get(key)),
                        csv_safe(value),
                        csv_safe(d["overrides"].get(key)),
                    ]
                )
        fields.freeze_panes = "A2"
        for col in "ABCDEF":
            fields.column_dimensions[col].width = 30
        stream = io.BytesIO()
        book.save(stream)
        return (
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    styles = getSampleStyleSheet()
    body = []

    def p(text, style="BodyText"):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

    body += [p(snapshot["report_title"], "Title"), Spacer(1, 12)]
    for label, value in [
        ("Case", snapshot["reference"]),
        ("Customer", snapshot["customer"]),
        ("Organization", snapshot["organization"]),
        ("Release", report.state),
        ("Version", report.version),
        ("SOP result", snapshot["validation_status"]),
        ("Configuration", snapshot["config_version"]),
        ("Provider", snapshot["provider"]),
        ("Approved by", report.approved_by or "Pending"),
    ]:
        body.append(p(f"{label}: {value}"))
    if snapshot["provider"] == "fixture":
        body += [
            Spacer(1, 12),
            p(
                "DEVELOPMENT FIXTURE — synthetic values; not extraction of document content.",
                "Heading2",
            ),
        ]
    if "scores" in snapshot:
        from .kertas_kerja import remarks

        body += [Spacer(1, 12), p("Kertas Kerja scores", "Heading2")]
        for key, value in snapshot["scores"].items():
            body.append(
                p(
                    f"{key.replace('_', ' ')}: {value if value is not None else 'Unresolved'}"
                )
            )
        body += [p(remarks(snapshot))]
    if snapshot.get("sop_summary"):
        points = snapshot["sop_summary"]
        body += [p("SOP rule points", "Heading2"), p(f"{points['earned']}/{points['possible']} points; {points['percentage']}%; {points['failed']} failed; {points['pending']} pending visual review. Provisional: {points['provisional']}. Release readiness is assessed separately.")]
    forex = snapshot.get("sop_packages", {}).get("forex")
    if forex:
        points = forex["summary"]
        body += [p("Forex: 15 SOP checkpoints", "Heading2"), p(f"{points['earned']}/{points['possible']} points; {points['percentage']}%; {points['failed']} failed; {points['pending']} pending visual review. All applicable child checks must pass for a checkpoint to earn one point.")]
        for checkpoint in forex["checkpoints"]:
            body.append(p(f"{checkpoint['rule_id']}: {checkpoint['result']} - {checkpoint.get('label','')}"))
    body += [Spacer(1, 18), p("Document inventory", "Heading2")]
    for d in snapshot["documents"]:
        body.append(p(f"{d['filename']} · {d['doc_type']} · {d['status']}"))
    body += [Spacer(1, 18), p("SOP findings", "Heading2")]
    for f in snapshot["findings"]:
        body += [
            p(f"{f['result']} · {f['field']} · {f['rule_id']}", "Heading3"),
            p(f.get("message", "")),
            p("Actual: " + json.dumps(f.get("actual"), ensure_ascii=False)),
            p("Expected: " + json.dumps(f.get("expected"), ensure_ascii=False)),
            p("Evidence: " + json.dumps(f.get("evidence"), ensure_ascii=False)),
        ]
    if not snapshot["findings"]:
        body.append(p("No findings recorded in this evaluation."))
    body += [Spacer(1, 18), p("Extracted values", "Heading2")]
    for d in snapshot["documents"]:
        body.append(p(d["filename"], "Heading3"))
        for key, value in d["normalized"].items():
            body.append(
                p(
                    key
                    + ": "
                    + (
                        value
                        if isinstance(value, str)
                        else json.dumps(value, ensure_ascii=False)
                    )
                )
            )
    from .storage import store
    from .field_evidence import evidence_images
    body += [Spacer(1, 18), p("Extracted field evidence", "Heading2")]
    for d in snapshot["documents"]:
        for e in evidence_images(d.get("validation", {})):
            body.append(p(f"{d['filename']} · {e.get('field', 'source region')} · page {e.get('page_number', 'unknown')} · observed {e.get('observed_present')} · image available {e['available']} · {e.get('review_status', 'PENDING')}"))
            if e["available"]:
                image = Image(io.BytesIO(store().read(e["storage_key"])))
                scale = min(400 / image.imageWidth, 160 / image.imageHeight, 1)
                image.drawWidth = image.imageWidth * scale
                image.drawHeight = image.imageHeight * scale
                body.append(image)
            else:
                body.append(p(e.get("error", "Bounding box unavailable; review required.")))
    stream = io.BytesIO()
    SimpleDocTemplate(stream, title=snapshot["report_title"], author="SQE").build(body)
    return stream.getvalue(), "application/pdf"


def report_key(report, fmt):
    from .storage import object_key

    return object_key(
        report.org,
        report.branch_id,
        "reports",
        report.id,
        f"{report.state.lower()}.{fmt}",
    )


def persist_report(report):
    """Store distinct draft/final artifacts; report IDs identify immutable snapshots."""
    from .storage import store

    formats = report.snapshot.get("configuration", {}).get("workflow", [])
    selected = next(
        (
            s["parameters"].get("formats")
            for s in formats
            if s["action"] == "render_report"
        ),
        ["json", "pdf", "xlsx"],
    )
    for fmt in dict.fromkeys(["json", *selected]):
        data, mime = make_report(report, fmt)
        store().put(report_key(report, fmt), data, mime)
