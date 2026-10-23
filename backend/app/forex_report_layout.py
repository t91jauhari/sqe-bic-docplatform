"""Excel presentation only. Never change extraction, findings or score calculation."""

from .configuration.seed_catalog import seed_value
import textwrap
from datetime import datetime
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CATEGORY_LABELS = seed_value("forex_report_layout.CATEGORY_LABELS")
CARRIER_LABELS = seed_value("forex_report_layout.CARRIER_LABELS")


def visual_review(finding, snapshot):
    if finding.get("evaluation_type"):
        return finding["evaluation_type"] == "VISUAL"
    from .sop_scoring import VISUAL_OPERATORS
    from .sop import fields

    rule = next(
        (
            r
            for r in snapshot.get("configuration", {}).get("sop_rules", [])
            if r["id"] == finding["rule_id"]
        ),
        {},
    )
    visual_paths = {
        f["path"]
        for f in fields()
        if f.get("type")
        in {
            "signature_region",
            "stamp_region",
            "photo_region",
            "page_marks",
            "page_observations",
        }
    }
    return (
        rule.get("operator") in VISUAL_OPERATORS or finding.get("field") in visual_paths
    )


def readable_remarks(snapshot, compact=False):
    groups = []
    for category, label in CATEGORY_LABELS.items():
        issues = [
            f
            for f in snapshot.get("findings", [])
            if f.get("eligible")
            and f.get("category") == category
            and f["result"] in {"FAIL", "REVIEW", "ERROR"}
        ]
        if not issues:
            continue
        lines = [f"=== {label} ==="]
        for n, finding in enumerate(issues[:2] if compact else issues, 1):
            message = " ".join(finding.get("message", "Check failed").split())
            if compact and len(message) > 190:
                message = message[:187] + "…"
            state = (
                (
                    "Awaiting visual review"
                    if visual_review(finding, snapshot)
                    else "Awaiting data assessment"
                )
                if finding["result"] == "REVIEW"
                else "Failed"
            )
            lines.append(f"{n}. {state}: {message}")
        if compact and len(issues) > 2:
            lines.append(f"+ {len(issues)-2} more findings; see SOP Details.")
        groups.append("\n".join(lines))
    if compact and groups:
        groups.append("Full findings, values and rule references: SOP Details.")
    if groups:
        return "\n\n".join(groups)
    applicable = [
        f
        for f in snapshot.get("findings", [])
        if f.get("eligible") and f["result"] != "NOT_APPLICABLE"
    ]
    return (
        "All applicable SOP checks passed."
        if applicable
        else "No applicable SOP checks assessed; see SOP Details for coverage."
    )


def display_score(snapshot, key, category=None):
    score = snapshot.get("scores", {}).get(key)
    if score is not None:
        return score
    categories = (
        {category}
        if category
        else {
            "total_a": {"DATA_COMPLETION", "CUSTOMER_SIGNATURE"},
            "total_b": {"DOCUMENT_COMPLETION", "OFFICER_REVIEW_APPROVAL"},
            "total": set(CATEGORY_LABELS),
        }.get(key, set())
    )
    relevant = [
        f
        for f in snapshot.get("findings", [])
        if f.get("eligible") and f.get("category") in categories
    ]
    if any(f["result"] == "REVIEW" and visual_review(f, snapshot) for f in relevant):
        return "Pending"
    if relevant and all(f["result"] == "NOT_APPLICABLE" for f in relevant):
        return "N/A"
    return "Not assessed"


def apply_layout(book, snapshots):
    if "Kertas Kerja" not in book:
        return
    snapshots = [
        s
        for s in snapshots
        if s.get("configuration", {})
        .get("templates", {})
        .get(s.get("template"), {})
        .get("report_layout")
        != "account_opening"
    ]
    if not snapshots:
        return
    sheet = book["Kertas Kerja"]
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 65
    thin = Side(style="thin", color="404040")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    grey = PatternFill("solid", fgColor="666666")
    for r in range(1, 8):
        sheet.cell(r, 3, ":")
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        sheet.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
        sheet.row_dimensions[r].height = 18
        sheet.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(r, 1).font = Font(name="Calibri", size=10, bold=True)
        sheet.cell(r, 4).font = Font(name="Calibri", size=10, bold=True)
    export = snapshots[0].get("report_export_context", {})
    if export.get("period"):
        sheet.cell(6, 4, export["period"])
    # Preserve configured business metadata; never invent a branch or audit period.
    for r, label in (
        (1, "Nama Kanwil"),
        (2, "Nama Cabang"),
        (3, "Category"),
        (4, "Aktivitas"),
        (5, "Jenis Pemeriksaan"),
        (6, "Periode"),
        (7, "Tanggal"),
    ):
        sheet.cell(r, 1, label)
    sheet.cell(
        8,
        1,
        "Indicators: 1 = compliant; 0 = failed; Pending = visual review; N/A = not applicable; Not assessed = no assessment. Rule points are on separate sheets.",
    )
    sheet.merge_cells("A8:T8")
    sheet.cell(8, 1).font = Font(name="Calibri", size=9, italic=True, color="555555")
    sheet.cell(8, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[8].height = 28
    headers = [
        "No",
        "Finding\nStatus",
        "Data\nCount",
        "Customer Name",
        "Account\nNo.",
        "Transaction\nType",
        "Inputter",
        "Approver",
        "Transaction\nCarrier",
        "Transaction\nDate",
        "Inspection\nDate",
    ]
    for col, label in enumerate(headers, 1):
        sheet.cell(10, col, label)
        sheet.merge_cells(start_row=10, start_column=col, end_row=11, end_column=col)
    for cell, label in (
        ("L10", "Total A"),
        ("O10", "Total B"),
        ("R10", "Total"),
        ("T10", "Remarks (Product)"),
    ):
        sheet[cell] = label
    for region in ("L10:N10", "O10:Q10", "R10:R11", "T10:T11"):
        sheet.merge_cells(region)
    for col, label in enumerate(
        [
            "Data\nCompletion",
            "Customer\nSignature",
            "Total A",
            "Document\nCompletion",
            "Officer Review &\nApproval",
            "Total B",
        ],
        12,
    ):
        sheet.cell(11, col, label)
    for row in sheet.iter_rows(min_row=10, max_row=11, max_col=20):
        for cell in row:
            cell.fill = grey
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
    sheet.row_dimensions[10].height = 20
    sheet.row_dimensions[11].height = 35
    widths = [
        6,
        12,
        8,
        36,
        17,
        17,
        16,
        16,
        20,
        16,
        16,
        13,
        13,
        12,
        14,
        17,
        12,
        12,
        8,
        115,
    ]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.column_dimensions["S"].hidden = True
    score_keys = [
        "DATA_COMPLETION",
        "CUSTOMER_SIGNATURE",
        "total_a",
        "DOCUMENT_COMPLETION",
        "OFFICER_REVIEW_APPROVAL",
        "total_b",
        "total",
    ]
    for n, snap in enumerate(snapshots, 12):
        for col in range(1, 21):
            cell = sheet.cell(n, col)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
                horizontal="left" if col in {4, 5, 7, 8, 9, 20} else "center",
            )
        for col in (7, 8):
            sheet.cell(n, col).fill = PatternFill("solid", fgColor="BFBFBF")
        carrier = snap.get("case_context", {}).get("transaction_carrier") or "UNKNOWN"
        sheet.cell(n, 9, CARRIER_LABELS.get(carrier, carrier))
        for col, key in enumerate(score_keys, 12):
            sheet.cell(
                n,
                col,
                display_score(snap, key, key if key in CATEGORY_LABELS else None),
            )
        remarks = readable_remarks(snap, compact=True)
        state = str(snap.get("release_state", "DRAFT")).replace("_", " ").title()
        pending = sum(
            f["result"] == "REVIEW" and visual_review(f, snap)
            for f in snap.get("findings", [])
            if f.get("eligible")
        )
        remarks += f"\n\nReport state: {state}. Awaiting visual review: {pending}."
        sheet.cell(n, 20, remarks)
        sheet.cell(n, 20).data_type = "s"
        sheet.cell(n, 20).alignment = Alignment(wrap_text=True, vertical="top")
        lines = sum(
            max(1, len(textwrap.wrap(line, 105))) for line in remarks.split("\n")
        )
        sheet.row_dimensions[n].height = min(409, max(110, lines * 11 + 18))
        sheet.cell(n, 5).number_format = "@"  # Account identifiers remain text.
        from .branch_report import transaction_date

        day = transaction_date(snap)
        if day:
            sheet.cell(n, 10, day).number_format = "dd-mmm-yy"
        try:
            inspection = datetime.strptime(
                str(sheet.cell(n, 11).value), "%d-%m-%Y"
            ).date()
            sheet.cell(n, 11, inspection).number_format = "dd-mmm-yy"
        except ValueError:
            pass
    summary = len(snapshots) + 13
    sheet.cell(summary, 1, "Total customers")
    for col, key in enumerate(score_keys, 12):
        values = [s.get("scores", {}).get(key) for s in snapshots]
        known = [v for v in values if v is not None]
        for offset, value in enumerate(
            [
                sum(known) if known else "Not assessed",
                100 * sum(known) / len(known) if known else "Not assessed",
                len(known),
                len(values) - len(known),
            ]
        ):
            cell = sheet.cell(summary + offset, col, value)
            cell.fill = PatternFill("solid", fgColor="E7E6E6")
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if offset == 1 and known:
                cell.number_format = '0.00"%"'
    for row in range(summary + 1, sheet.max_row + 1):
        sheet.cell(row, 4).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = max(
            sheet.row_dimensions[row].height or 15, 30
        )
    sheet.freeze_panes = "F12"
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.oddFooter.right.text = "Branch Kertas Kerja"
