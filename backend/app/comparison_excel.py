"""Styled application export of authorized, current comparison snapshots."""
import io
import json
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def cell_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, sort_keys=True)
    # Store user/OCR text literally, including formula prefixes.
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def export_workbook(results):
    book = Workbook()
    book.remove(book.active)
    summary = book.create_sheet("Summary")
    documents = book.create_sheet("Document results")
    fields = book.create_sheet("Field comparisons")
    edits = book.create_sheet("Reference edits")
    about = book.create_sheet("Export information")
    summary.append(["Customer", "Case ID", "Sample", "Revision", "Result", "Match rate (%)", "Complete coverage", "Compared fields", "Match", "Mismatch", "Missing value", "Missing document", "Review", "Unmapped", "Pending", "Not applicable", "Fingerprint"])
    fields.append(["Customer", "Case ID", "Sample", "Revision", "Document", "Filename", "Document ID", "Field", "Field label", "Original reference", "Current reference", "Effective normalized value", "Result", "Explanation", "Reference edited", "Extraction ID", "Fingerprint"])
    edits.append(["Customer", "Case ID", "Sample", "Field", "Field label", "Original reference", "Edited reference", "Reason", "Actor", "Edited at (UTC)"])
    documents.append(["Customer", "Case ID", "Sample", "Document", "Filename", "Document ID", "Result", "Match rate (%)", "Exact match rate (%)", "Coverage (%)", "Compared fields", "Applicable fields", "Exact matches", "Fuzzy matches", "Review", "Unmapped", "Pending"])
    summary.cell(1, 18, "Fuzzy matches")
    fields.cell(1, 18, "Similarity (%)")
    fields.cell(1, 19, "Comparison method")
    for result in results:
        stats = result["summary"]
        counts = stats.get("counts", {})
        prefix = [result["customer"], result["case_id"], result["sample_name"], result["revision"]]
        summary.append([*prefix, result["result"], stats.get("match_rate"), stats.get("complete"), stats.get("compared_fields"), *[counts.get(k, 0) for k in ("MATCH", "MISMATCH", "VALUE_MISSING", "DOCUMENT_MISSING", "REVIEW", "UNMAPPED", "PENDING", "NOT_APPLICABLE")], result["fingerprint"]])
        summary.cell(summary.max_row, 18, counts.get("FUZZY_MATCH", 0))
        for group in result.get("document_groups", []):
            gc = group["counts"]
            documents.append([*prefix[:3], group["document_label"], group.get("filename"), group.get("document_id"), group["result"], group["match_rate"], group["exact_match_rate"], group["coverage_rate"], group["compared_fields"], group["applicable_fields"], *[gc.get(k, 0) for k in ("MATCH", "FUZZY_MATCH", "REVIEW", "UNMAPPED", "PENDING")]])
        for row in result["rows"]:
            fields.append([cell_value(v) for v in [*prefix, row["document_label"], row.get("filename"), row.get("document_id"), row.get("field"), row["label"], row["original_value"], row["dataset_value"], row.get("effective_value"), row["status"], row["message"], row["edited"], row.get("extraction_id"), result["fingerprint"]]])
            fields.cell(fields.max_row, 18, row.get("similarity"))
            fields.cell(fields.max_row, 19, row.get("method", "exact_or_strict"))
            edit = row.get("edit")
            if edit:
                edits.append([cell_value(v) for v in [*prefix[:3], row.get("field"), row["label"], row["original_value"], row["dataset_value"], edit["reason"], edit["actor"], datetime.fromtimestamp(edit["at"], timezone.utc).isoformat()]])
    about.append(["Item", "Value"])
    about.append(["Exported at (UTC)", datetime.now(timezone.utc).isoformat()])
    about.append(["Comparisons", len(results)])
    about.append(["Scope", "All saved comparisons accessible to the signed-in user's organization and branches; merged cases excluded."])
    about.append(["Values", "Latest saved references and current effective normalized values. Unsaved browser edits are excluded; pending revalidation is labelled PENDING."])
    about.append(["Match rate", "(MATCH + FUZZY_MATCH) / (MATCH + FUZZY_MATCH + MISMATCH + VALUE_MISSING). Exact matching and applicable-field coverage are separate in Document results."])
    about.append(["Policy", results[0]["policy"] if results else "No saved comparisons."])
    about.append(["Dataset SHA-256", results[0].get("dataset_sha256") if results else None])
    colors = {"MATCH": "EAF5EF", "FUZZY_MATCH": "E7F2F7", "MISMATCH": "FBEDED", "VALUE_MISSING": "FBEDED", "DOCUMENT_MISSING": "FBEDED", "REVIEW": "FFF4DC", "UNMAPPED": "FFF4DC", "PENDING": "EEF2F6"}
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        for row in sheet:
            for cell in row:
                cell.value = cell_value(cell.value)
                cell.font = Font(name="Calibri", size=11, color="243B45")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="167B70")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F7F7")
                if cell.row > 1 and isinstance(cell.value, str) and cell.value in colors:
                    cell.fill = PatternFill("solid", fgColor=colors[cell.value])
                    cell.font = Font(name="Calibri", size=11, bold=True, color="243B45")
                cell.border = Border(bottom=Side(style="hair", color="DCE6E8"))
        for col in range(1, sheet.max_column + 1):
            header = str(sheet.cell(1, col).value)
            width = 44 if header in {"Original reference", "Current reference", "Effective normalized value", "Edited reference", "Explanation", "Reason"} else 32 if "ID" in header or header in {"Fingerprint", "Filename", "Field label", "Document"} else 23
            sheet.column_dimensions[get_column_letter(col)].width = width
        for row in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 60 if sheet == fields or sheet == edits else 32
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    about.column_dimensions["B"].width = 110
    for row in range(2, about.max_row + 1):
        about.row_dimensions[row].height = 60
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
