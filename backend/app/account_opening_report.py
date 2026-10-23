"""Account Opening Kertas Kerja layout, preserving evidence sheets and machine snapshots."""

from .configuration.seed_catalog import seed_value

import copy
import io
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .sop import value_at
from datetime import datetime
from zoneinfo import ZoneInfo

HEADERS = seed_value("account_opening_report.HEADERS")


def render(snapshots, legacy):
    from .kertas_kerja import excel_value

    book = load_workbook(io.BytesIO(legacy(snapshots)))
    sheet = book.worksheets[0]
    sheet.title = "Account Opening Kertas Kerja"
    sheet.delete_rows(11, max(sheet.max_row - 10, 1))
    for col, label in enumerate(HEADERS, 1):
        sheet.cell(11, col, label)
    for number, snap in enumerate(snapshots, 1):
        from .configuration.reporting import project_report

        projected, timezone = project_report(snap, "account_opening")
        scores = snap.get("scores", {})
        remarks = []
        for finding in snap["findings"]:
            if finding["result"] in {"FAIL", "REVIEW", "ERROR"}:
                remarks.append(
                    f"{finding['category']} · {finding['rule_id']}: {finding['result']} — {finding['message']}"
                )
        remarks += ["Required document missing: " + v for v in snap.get("missing", [])]
        # The supplied schema labels this as an existing account, not the new account.
        if not projected["account"]:
            remarks.append(
                "New account number is not a separate field in the supplied extraction specification; provide document-supported report context rather than copying an existing account number."
            )
        values = [
            number,
            (
                "Yes"
                if scores.get("total") == 0
                else "Review" if scores.get("total") is None else "No"
            ),
            1,
            projected["customer"],
            projected["account"],
            projected["transaction_type"],
            projected["inputter"],
            projected["approver"],
            projected["carrier"],
            projected["transaction_date"],
            datetime.fromtimestamp(
                snap.get("sop_completed_at", snap["generated_at"]),
                ZoneInfo(timezone),
            ).strftime("%d-%m-%Y"),
            scores.get("DATA_COMPLETION"),
            scores.get("CUSTOMER_SIGNATURE"),
            scores.get("OFFICER_REVIEW_APPROVAL"),
            scores.get("total_a"),
            scores.get("DOCUMENT_COMPLETION"),
            scores.get("DOCUMENT_VERIFICATION"),
            scores.get("OBSERVATIONS"),
            scores.get("total_b"),
            scores.get("total"),
            "\n".join(remarks),
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(number + 11, col, excel_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[number + 11].height = 110
    for cell in sheet[11]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "D12"
    sheet.auto_filter.ref = f"A11:U{11 + len(snapshots)}"
    for col in ["L", "M", "N", "O", "P", "Q", "R", "S", "T"]:
        sheet.column_dimensions[col].width = 18
    sheet.column_dimensions["S"].hidden = False
    sheet.column_dimensions["U"].width = 85
    end = 14 + len(snapshots)
    sheet.cell(end, 1, "Total cases")
    sheet.cell(end, 3, len(snapshots))
    for col, key in [(15, "total_a"), (19, "total_b"), (20, "total")]:
        values = [v.get("scores", {}).get(key) for v in snapshots]
        known = [v for v in values if v is not None]
        sheet.cell(end, col, sum(known) if known else None)
        sheet.cell(
            end + 1, col, sum(known) / len(known) if known else None
        ).number_format = "0.00%"
    sheet.cell(
        end + 1, 4, "Scored compliance; unresolved cases excluded from denominator"
    )
    sheet.print_area = f"A1:U{end + 3}"
    for col, key in [(15, "total_a"), (19, "total_b"), (20, "total")]:
        values = [v.get("scores", {}).get(key) for v in snapshots]
        sheet.cell(end + 2, col, sum(v is not None for v in values))
        sheet.cell(end + 3, col, sum(v is None for v in values))
    sheet.cell(end + 2, 4, "Scored cases")
    sheet.cell(end + 3, 4, "Unresolved cases")
    from .report_summary import write_summary

    summary_end = write_summary(sheet, end + 1, snapshots)
    sheet.print_area = f"A1:U{summary_end}"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def workbook(snapshots, legacy):
    account = [
        s
        for s in snapshots
        if s.get("configuration", {})
        .get("templates", {})
        .get(s.get("template"), {})
        .get("report_layout")
        == "account_opening"
    ]
    other = [s for s in snapshots if s not in account]
    if not account:
        return legacy(snapshots)
    if not other:
        return render(account, legacy)
    result = load_workbook(io.BytesIO(legacy(other)))
    addition = load_workbook(io.BytesIO(render(account, legacy)))
    for source in addition:
        target = result.create_sheet(("AO " + source.title)[:31])
        for row in source:
            for cell in row:
                out = target.cell(cell.row, cell.column, cell.value)
                if cell.has_style:
                    out.font = copy.copy(cell.font)
                    out.fill = copy.copy(cell.fill)
                    out.border = copy.copy(cell.border)
                    out.number_format = cell.number_format
                    out.protection = copy.copy(cell.protection)
                out.alignment = copy.copy(cell.alignment)
        for key, dim in source.column_dimensions.items():
            target.column_dimensions[key] = copy.copy(dim)
        for key, dim in source.row_dimensions.items():
            target.row_dimensions[key] = copy.copy(dim)
        for merged in source.merged_cells.ranges:
            target.merge_cells(str(merged))
        target.freeze_panes = source.freeze_panes
        target.auto_filter = copy.copy(source.auto_filter)
        target.page_setup = copy.copy(source.page_setup)
        target.print_options = copy.copy(source.print_options)
        target.sheet_properties = copy.copy(source.sheet_properties)
        if source.print_area:
            target.print_area = source.print_area
        # Keep source evidence image objects as independent copied objects.
        for visual in source._images:
            target.add_image(copy.copy(visual))
    stream = io.BytesIO()
    result.save(stream)
    return stream.getvalue()
