"""Server-computed Kertas Kerja scores and deterministic remarks, one row per case."""

import io
import json
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .sop import CATEGORIES, value_at
from .value_display import display_value

HEADERS = [
    "No.",
    "Finding Status",
    "Data Count",
    "Customer Name",
    "Account No.",
    "Transaction Type",
    "Inputter",
    "Approver",
    "Transaction Carrier",
    "Transaction Date",
    "Inspection Date",
    "Data Completion",
    "Customer Signature",
    "Total A",
    "Document Completion",
    "Officer Review & Approval",
    "Total B",
    "Total",
    "BIC Remarks",
    "Product Remarks",
]


def text(value):
    if value is None:
        return ""
    return display_value(value)


def excel_value(value):
    """Format compound observations without mutating machine results."""
    if isinstance(value, (dict, list, tuple)):
        value = display_value(list(value) if isinstance(value, tuple) else value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def remarks(snapshot):
    from .forex_report_layout import readable_remarks
    return readable_remarks(snapshot)


def row(snapshot, n):
    from .configuration.reporting import project_report
    projected, timezone = project_report(snapshot, 'forex')
    scores = snapshot.get("scores", {})

    results = [f["result"] for f in snapshot["findings"] if f.get("eligible")]
    finding = (
        "Yes"
        if "FAIL" in results
        else "Review"
        if any(x in {"REVIEW", "ERROR"} for x in results)
        or snapshot["validation_status"] == "ERROR"
        else "No"
    )
    return [
        n,
        finding,
        1,
        projected['customer'],
        projected['account'],
        projected['transaction_type'],
        projected['inputter'],
        projected['approver'],
        projected['carrier'],
        projected['transaction_date'],
        datetime.fromtimestamp(
            snapshot.get("sop_completed_at", snapshot["generated_at"]),
            ZoneInfo(timezone),
        ).strftime("%d-%m-%Y"),
        scores.get(CATEGORIES[0]),
        scores.get(CATEGORIES[1]),
        scores.get("total_a"),
        scores.get(CATEGORIES[2]),
        scores.get(CATEGORIES[3]),
        scores.get("total_b"),
        scores.get("total"),
        None,
        remarks(snapshot),
    ]


def _legacy_workbook(snapshots):
    w = Workbook()
    s = w.active
    s.title = "Kertas Kerja"
    first = snapshots[0]
    ctx = first.get("configuration", {}).get("report_context", {})
    for n, (label, v) in enumerate(
        [
            ("Nama Kanwil", ctx.get("region_name", "")),
            (
                "Nama Cabang",
                first.get("branch_code", "") + " " + first.get("branch_name", ""),
            ),
            ("Category", ctx.get("category", "Forex")),
            ("Aktivitas", ctx.get("activity", "Document review")),
            ("Jenis Pemeriksaan", ctx.get("inspection_type", "SOP")),
            ("Periode", ctx.get("period", "")),
            (
                "Tanggal Export",
                datetime.now(ZoneInfo("Asia/Jakarta")).strftime("%d-%m-%Y %H:%M %Z"),
            ),
        ],
        1,
    ):
        s.cell(n, 1, label)
        s.cell(n, 4, v)
    if any(snap["provider"] == "fixture" for snap in snapshots):
        s.cell(
            9, 1, "DEVELOPMENT FIXTURE — synthetic extraction; not bank audit evidence."
        )
    for n, header in enumerate(HEADERS, 1):
        s.cell(11, n, header)
    for n, snapshot in enumerate(snapshots, 1):
        for col, v in enumerate(row(snapshot, n), 1):
            s.cell(n + 11, col, excel_value(v))
        s.row_dimensions[n + 11].height = 110
    end = 11 + len(snapshots)
    summary = end + 2
    s.cell(summary, 1, "Total cases")
    s.cell(summary, 3, len(snapshots))
    for col, key in [(14, "total_a"), (17, "total_b"), (18, "total")]:
        values = [x.get("scores", {}).get(key) for x in snapshots]
        known = [v for v in values if v is not None]
        s.cell(summary, col, sum(known) if known else None)
        s.cell(summary + 1, 4, "Scored compliance (%)")
        s.cell(summary + 1, col, 100 * sum(known) / len(known) if known else None)
        s.cell(summary + 2, 4, "Scored cases")
        s.cell(summary + 2, col, len(known))
        s.cell(summary + 3, 4, "Unresolved cases")
        s.cell(summary + 3, col, len(values) - len(known))
    from .report_summary import write_summary
    summary_end = write_summary(s, summary + 1, snapshots)
    s.column_dimensions["S"].hidden = True
    from openpyxl.utils import get_column_letter

    for col in range(1, 21):
        s.column_dimensions[get_column_letter(col)].width = (
            18 if col not in {4, 20} else 45
        )
    s.freeze_panes = "L12"
    s.auto_filter.ref = f"A11:T{end}"
    s.print_title_rows = "1:11"
    s.print_area = f"A1:T{summary_end}"
    s.page_setup.orientation = "landscape"
    s.page_setup.paperSize = s.PAPERSIZE_A3
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
    evidence_sheet = w.create_sheet("Visual Evidence")
    evidence_sheet.append(
        [
            "Case",
            "Document",
            "Field",
            "Page",
            "Observed presence",
            "Image available",
            "Review",
            "Source SHA256",
            "Revision",
            "BBox",
            "Image / error",
        ]
    )
    from openpyxl.drawing.image import Image as SheetImage
    from .storage import store

    for snap in snapshots:
        for d in snap["documents"]:
            for e in d.get("validation", {}).get("visual_evidence", []):
                evidence_sheet.append(
                    [
                        snap["reference"],
                        d["filename"],
                        e["field"],
                        e.get("page_number"),
                        text(e.get("observed_present")),
                        e["available"],
                        e["review_status"],
                        e["source_sha256"],
                        e["revision"],
                        text(e.get("bbox")),
                        e.get("error", ""),
                    ]
                )
                n = evidence_sheet.max_row
                if e["available"]:
                    image = SheetImage(io.BytesIO(store().read(e["storage_key"])))
                    scale = min(280 / image.width, 150 / image.height, 1)
                    image.width *= scale
                    image.height *= scale
                    evidence_sheet.add_image(image, f"K{n}")
                    evidence_sheet.row_dimensions[n].height = 120
    evidence_sheet.column_dimensions["K"].width = 44
    evidence_sheet.freeze_panes = "A2"
    details = w.create_sheet("SOP Details")
    details.append(
        [
            "Case",
            "Rule",
            "Source",
            "Scope",
            "Category",
            "Result",
            "Actual",
            "Expected",
            "Message",
            "Evidence",
            "Config version",
            "Report state",
        ]
    )
    quality = w.create_sheet("Document Quality")
    quality.append(["Case", "Document field", "Check", "Result", "Evidence", "Scoring"])
    values = w.create_sheet("Extracted Fields")
    values.append(
        ["Case", "Document", "Type", "Field", "Raw", "Normalized", "Override", "SHA256"]
    )
    for snap in snapshots:
        for f in snap["findings"]:
            details.append(
                [
                    snap["reference"],
                    f["rule_id"],
                    f.get("source_ref"),
                    f.get("scope_basis"),
                    f.get("category"),
                    f["result"],
                    text(f.get("actual")),
                    text(f.get("expected")),
                    f["message"],
                    text(f.get("evidence")),
                    snap["config_version"],
                    snap.get("release_state", "DRAFT"),
                ]
            )
        for f in snap.get("technical_findings", []):
            quality.append(
                [
                    snap["reference"],
                    f["field"],
                    f["rule_id"],
                    f["result"],
                    text(f.get("evidence")),
                    "Excluded from SOP scoring",
                ]
            )
        for d in snap["documents"]:
            for key, val in d["normalized"].items():
                values.append(
                    [
                        snap["reference"],
                        d["filename"],
                        d["doc_type"],
                        key,
                        text(d["raw"].get(key)),
                        text(val),
                        text(d["overrides"].get(key)),
                        d["sha256"],
                    ]
                )
    trace = w.create_sheet("Report Trace")
    trace.append(
        ["Case", "Configuration", "Generated", "Snapshot Hash", "Provider", "Context"]
    )
    for snap in snapshots:
        trace.append(
            [
                snap["reference"],
                snap["config_version"],
                snap["generated_at"],
                snap.get("snapshot_hash"),
                snap["provider"],
                text(snap.get("case_context", {})),
            ]
        )
    for sheet in w:
        header = 11 if sheet == s else 1
        for c in sheet[header]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="17344A")
        for r in sheet.iter_rows():
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, str):
                    c.data_type = "s"
        if sheet != s:
            sheet.freeze_panes = "A2"
            for i in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(i)].width = 30
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def workbook(snapshots):
    from .account_opening_report import workbook as render

    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(render(snapshots, _legacy_workbook)))
    from .forex_report_layout import apply_layout
    apply_layout(book, snapshots)
    summary = book.create_sheet("SOP Rule Points")
    summary.append(["Customer", "Case", "Earned", "Possible", "Percentage", "Failed", "Pending visual review", "Resolved", "Provisional"])
    details = book.create_sheet("SOP Point Detail")
    details.append(["Customer", "Rule", "Result", "Points", "Possible", "Evaluation type", "Reason"])
    package = book.create_sheet("Forex SOP Package")
    package.append(["Customer", "Checkpoint", "Result", "Points", "Possible", "Field checks", "Package coverage", "Description"])
    for snap in snapshots:
        forex = snap.get("sop_packages", {}).get("forex")
        if forex:
            for checkpoint in forex["checkpoints"]:
                applicable = checkpoint["result"] != "NOT_APPLICABLE"
                package.append([text(snap.get("customer")),checkpoint["rule_id"],checkpoint["result"],int(checkpoint["result"]=="PASS"),int(applicable),", ".join(checkpoint["child_rules"]),f"{forex['coverage']['present']}/15",checkpoint.get("label","")])
        points = snap.get("sop_summary")
        if not points:
            continue  # Historical snapshots retain their original contract.
        summary.append([text(snap.get("customer")), text(snap.get("reference")), *[points.get(k) for k in ("earned", "possible", "percentage", "failed", "pending", "resolved", "provisional")]])
        for f in snap.get("findings", []):
            eligible = f.get("eligible", True) and f["result"] != "NOT_APPLICABLE"
            details.append([text(snap.get("customer")), text(f["rule_id"]), f["result"], int(eligible and f["result"] == "PASS"), int(eligible), f.get("evaluation_type"), f.get("reason_code")])
    for sheet in (summary, details, package):
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 25
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
