"""Adjacent, explicit report totals; no formulas or ambiguous blank counts."""


def write_summary(sheet, row, snapshots):
    legacy = [snap.get("scores", {}).get("total") for snap in snapshots]
    known = [value for value in legacy if value is not None]
    rows = [
        ("Scored compliance (%) · legacy", 100 * sum(known) / len(known) if known else "Awaiting assessment"),
        ("Scored cases · legacy", len(known)),
        ("Unresolved cases · legacy", len(legacy) - len(known)),
    ]
    points = [snap["sop_summary"] for snap in snapshots if snap.get("sop_summary")]
    possible = sum(p["possible"] for p in points)
    earned = sum(p["earned"] for p in points)
    pending = sum(p["pending"] for p in points)
    rows += [
        ("Rule compliance (%)", round(100 * earned / possible, 2) if possible else "No applicable scored rules"),
        ("Rule points earned", earned),
        ("Rule points possible", possible),
        ("Failed rules", sum(p["failed"] for p in points)),
        ("Pending visual reviews", pending),
        ("Cases with rule-point scores", len(points)),
        ("Cases awaiting visual review", sum(p["pending"] > 0 for p in points)),
        ("Cases needing current SOP scoring", len(snapshots) - len(points)),
        ("Rule scoring status", "Provisional" if pending else "Incomplete coverage" if len(points) != len(snapshots) else "Complete" if possible else "No applicable rules"),
    ]
    for index, (label, value) in enumerate(rows):
        sheet.cell(row + index, 4, label)
        cell = sheet.cell(row + index, 5, value)
        if isinstance(value, (int, float)) and "(%)" in label:
            cell.number_format = '0.00"%"'
    sheet.cell(row + len(rows), 4, "Rule percentage includes pending visual checks in possible points. Legacy unresolved counts reflect binary totals; they are separate from pending visual reviews.")
    sheet.row_dimensions[row + len(rows)].height = 60
    from openpyxl.styles import Alignment
    sheet.cell(row + len(rows), 4).alignment = Alignment(wrap_text=True, vertical="top")
    return row + len(rows)
